from docx import Document
from docx.shared import Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn

doc = Document()

doc.styles['Normal'].font.name = '微软雅黑'
doc.styles['Normal']._element.rPr.rFonts.set(qn('w:eastAsia'), '微软雅黑')
doc.styles['Normal'].font.size = Pt(11)

# 标题
title = doc.add_paragraph()
title.alignment = WD_ALIGN_PARAGRAPH.CENTER
run = title.add_run('个人简历')
run.font.size = Pt(22)
run.font.bold = True
run.font.color.rgb = RGBColor(0x1A, 0x1A, 0x2E)

info = doc.add_paragraph()
info.alignment = WD_ALIGN_PARAGRAPH.CENTER
info_run = info.add_run('求职意向：全栈开发工程师（偏前端） | 2026年6月毕业 | 河北科技学院 · 软件工程')
info_run.font.size = Pt(11)
info_run.font.color.rgb = RGBColor(0x55, 0x55, 0x55)

doc.add_paragraph()

# 教育背景
h2 = doc.add_paragraph()
h2_run = h2.add_run('教育背景')
h2_run.font.size = Pt(14)
h2_run.font.bold = True
h2_run.font.color.rgb = RGBColor(0x00, 0x7A, 0xFF)

edu = doc.add_paragraph()
edu.add_run('河北科技学院  |  软件工程（本科）  |  2022.09 - 2026.06').font.size = Pt(11)

doc.add_paragraph()

# 专业技能
h2 = doc.add_paragraph()
h2_run = h2.add_run('专业技能')
h2_run.font.size = Pt(14)
h2_run.font.bold = True
h2_run.font.color.rgb = RGBColor(0x00, 0x7A, 0xFF)

skills = [
    '前端开发：熟练 HTML5 / CSS3 / JavaScript，能独立开发单页应用，掌握响应式布局与现代CSS特性',
    'UI设计：熟悉 iOS 风格毛玻璃设计体系，注重交互细节与用户体验，能独立完成界面设计',
    'AI辅助开发：熟练运用 AI 工具（豆包/ChatGPT/Cursor等）进行代码生成、调试优化、需求拆解，开发效率提升显著',
    '后端基础：了解 Java 与 Spring Boot 基本原理，熟悉 RESTful API 设计，能对接后端接口',
    '工程化：熟练使用 Git 版本管理，有 GitHub 项目托管与 Pages 部署经验，具备完整上线能力',
    '存储技术：了解 localStorage、File System Access API、IndexedDB 等浏览器本地存储方案',
]
for s in skills:
    p = doc.add_paragraph(style='List Bullet')
    p.add_run(s).font.size = Pt(11)

doc.add_paragraph()

# 项目经历
h2 = doc.add_paragraph()
h2_run = h2.add_run('项目经历')
h2_run.font.size = Pt(14)
h2_run.font.bold = True
h2_run.font.color.rgb = RGBColor(0x00, 0x7A, 0xFF)

proj_title = doc.add_paragraph()
proj_title.add_run('个人工作台 — 本地工作生活管理工具').font.bold = True
proj_title.runs[0].font.size = Pt(12)

link_p = doc.add_paragraph()
link_run = link_p.add_run('在线体验：https://kyoume819.github.io/my-workbench/')
link_run.font.size = Pt(10)
link_run.font.color.rgb = RGBColor(0x00, 0x66, 0xCC)

link_p2 = doc.add_paragraph()
link_run2 = link_p2.add_run('代码仓库：https://github.com/kyoume819/my-workbench')
link_run2.font.size = Pt(10)
link_run2.font.color.rgb = RGBColor(0x00, 0x66, 0xCC)

desc = doc.add_paragraph()
desc.add_run('项目描述：').font.bold = True
desc.add_run('独立开发的单文件HTML个人工作台应用，整合9大功能模块，纯前端实现无需服务器。全程使用AI辅助开发，从需求到上线效率大幅提升。已部署GitHub Pages，手机电脑均可访问。')

tech = doc.add_paragraph()
tech.add_run('技术栈：').font.bold = True
tech.add_run('HTML5 / CSS3 / 原生JavaScript / iOS毛玻璃UI / AI辅助开发 / Git / GitHub Pages')

work = doc.add_paragraph()
work.add_run('主要工作：').font.bold = True

works = [
    '前端架构：设计左侧导航+右侧内容区的单页应用架构，9个模块无刷新切换，状态管理清晰',
    'UI设计：采用iOS风格毛玻璃设计体系，半透明磨砂质感，响应式适配桌面与移动端，视觉效果出众',
    '功能开发：独立完成今日计划、自媒体追踪、项目管理、健身记录、饮食营养计算、攒钱目标等全部业务模块',
    'AI辅助：全程运用AI工具辅助开发，包括代码生成、Bug调试、样式优化、需求拆解，开发效率提升2倍以上',
    '数据持久化：实现本地数据自动保存与备份恢复，兼容File System Access API与localStorage双方案',
    '工程化：Git版本管理12次迭代提交，部署GitHub Pages生成在线链接，具备完整从开发到上线的全流程能力'
]
for w in works:
    p = doc.add_paragraph(style='List Bullet')
    p.add_run(w).font.size = Pt(11)

doc.add_paragraph()

# 自我评价
h2 = doc.add_paragraph()
h2_run = h2.add_run('自我评价')
h2_run.font.size = Pt(14)
h2_run.font.bold = True
h2_run.font.color.rgb = RGBColor(0x00, 0x7A, 0xFF)

evals = [
    '前端开发能力扎实，能独立完成从界面设计到代码实现的全流程，注重用户体验与视觉细节',
    '熟练运用AI工具提升开发效率，善于借助AI解决技术问题，快速学习新技术',
    '有完整项目从0到1的独立开发经验，熟悉软件开发全流程',
    '学习能力强，做事踏实细心，对前端技术有热情，喜欢动手做实用的产品'
]
for e in evals:
    p = doc.add_paragraph(style='List Bullet')
    p.add_run(e).font.size = Pt(11)

doc.save(r'C:\Users\a1773\Desktop\kk的APP\简历-全栈偏前端方向.docx')
print('简历生成完成')
