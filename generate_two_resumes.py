from docx import Document
from docx.shared import Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn

def create_resume(filename, target_title, skills, project_desc, works, evals):
    doc = Document()
    doc.styles['Normal'].font.name = '微软雅黑'
    doc.styles['Normal']._element.rPr.rFonts.set(qn('w:eastAsia'), '微软雅黑')
    doc.styles['Normal'].font.size = Pt(11)

    # 标题
    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title.add_run('个人简历')
    run.font.size = Pt(24)
    run.font.bold = True
    run.font.color.rgb = RGBColor(0x1A, 0x1A, 0x2E)

    info = doc.add_paragraph()
    info.alignment = WD_ALIGN_PARAGRAPH.CENTER
    info_run = info.add_run(target_title)
    info_run.font.size = Pt(11)
    info_run.font.color.rgb = RGBColor(0x55, 0x55, 0x55)

    doc.add_paragraph()

    # 专业技能（放前面）
    h2 = doc.add_paragraph()
    h2_run = h2.add_run('专业技能')
    h2_run.font.size = Pt(14)
    h2_run.font.bold = True
    h2_run.font.color.rgb = RGBColor(0x00, 0x7A, 0xFF)
    h2_run.space_after = Pt(6)

    for s in skills:
        p = doc.add_paragraph(style='List Bullet')
        p.add_run(s).font.size = Pt(11)

    doc.add_paragraph()

    # 项目经历（重点）
    h2 = doc.add_paragraph()
    h2_run = h2.add_run('项目经历')
    h2_run.font.size = Pt(14)
    h2_run.font.bold = True
    h2_run.font.color.rgb = RGBColor(0x00, 0x7A, 0xFF)

    proj_title = doc.add_paragraph()
    proj_title.add_run('个人工作台 — 本地工作生活管理工具').font.bold = True
    proj_title.runs[0].font.size = Pt(12)

    link_p = doc.add_paragraph()
    link_run = link_p.add_run('在线体验：https://kyoume819.github.io/my-workbench/')
    link_run.font.size = Pt(10)
    link_run.font.color.rgb = RGBColor(0x00, 0x66, 0xCC)

    link_p2 = doc.add_paragraph()
    link_run2 = link_p2.add_run('代码仓库：https://github.com/kyoume819/my-workbench')
    link_run2.font.size = Pt(10)
    link_run2.font.color.rgb = RGBColor(0x00, 0x66, 0xCC)

    desc = doc.add_paragraph()
    desc.add_run('项目描述：').font.bold = True
    desc.add_run(project_desc)

    work = doc.add_paragraph()
    work.add_run('主要工作：').font.bold = True
    for w in works:
        p = doc.add_paragraph(style='List Bullet')
        p.add_run(w).font.size = Pt(11)

    doc.add_paragraph()

    # 自我评价
    h2 = doc.add_paragraph()
    h2_run = h2.add_run('自我评价')
    h2_run.font.size = Pt(14)
    h2_run.font.bold = True
    h2_run.font.color.rgb = RGBColor(0x00, 0x7A, 0xFF)

    for e in evals:
        p = doc.add_paragraph(style='List Bullet')
        p.add_run(e).font.size = Pt(11)

    doc.add_paragraph()

    # 教育背景（放最后，不突出）
    h2 = doc.add_paragraph()
    h2_run = h2.add_run('教育背景')
    h2_run.font.size = Pt(12)
    h2_run.font.bold = True
    h2_run.font.color.rgb = RGBColor(0x88, 0x88, 0x88)

    edu = doc.add_paragraph()
    edu_run = edu.add_run('河北科技学院  |  软件工程（本科）  |  2022.09 - 2026.06')
    edu_run.font.size = Pt(10.5)
    edu_run.font.color.rgb = RGBColor(0x66, 0x66, 0x66)

    doc.save(filename)
    print(f'{filename} 生成完成')


# ===== 版本1：项目经理助理方向 =====
pm_skills = [
    '项目管理：熟悉软件开发全流程，能独立编写PRD、开发计划、进度追踪文档',
    '任务管理：熟练拆解任务、跟进进度、协调资源，推动项目按计划交付',
    '文档能力：能独立撰写需求文档、项目总结、会议纪要等各类项目文档',
    '沟通协调：具备良好的沟通能力和团队协作意识，能对接开发、设计、测试等不同角色',
    '技术背景：了解前端开发基本流程，能与开发团队有效沟通，理解技术术语和需求',
    '工具：熟练使用 Office 办公软件，了解 Git 版本管理，有项目部署上线经验'
]

pm_project_desc = '作为项目负责人，独立完成从需求分析、PRD撰写、开发计划制定到最终上线的全流程管理。项目包含9个功能模块，覆盖工作、自媒体、健身、饮食、财务等多个场景，已部署上线并持续迭代优化。'

pm_works = [
    '需求管理：独立撰写产品需求文档（PRD），明确产品目标、功能边界、验收标准，完成需求拆解与优先级排序',
    '进度管理：制定分阶段开发计划，将项目拆分为12个阶段任务，跟踪每个阶段完成情况，确保按计划交付',
    '文档管理：编写并维护PRD文档、开发计划、设计原型说明等各类项目文档，版本清晰、可追溯',
    '迭代管理：经历4轮需求迭代优化，每轮收集反馈、调整优先级、跟进落地，持续改进产品体验',
    '上线部署：负责项目版本管理与部署上线，使用Git进行版本控制，通过GitHub Pages发布',
    '质量把控：对照验收标准逐条验证功能，发现问题及时修复，确保交付质量'
]

pm_evals = [
    '有完整项目从0到1的管理经验，熟悉软件开发全流程，能快速理解需求并推进落地',
    '做事细心有条理，擅长文档撰写和进度跟踪，能把复杂任务拆解成可执行的小步骤',
    '沟通能力好，有技术背景，能和开发、设计、测试等不同角色有效对接',
    '执行力强，有责任心，善于协调资源解决问题，推动项目按计划完成'
]

create_resume(
    r'C:\Users\a1773\Desktop\kk的APP\简历-项目经理助理方向.docx',
    '求职意向：项目经理助理 | 2026年6月毕业',
    pm_skills,
    pm_project_desc,
    pm_works,
    pm_evals
)


# ===== 版本2：全栈偏前端方向 =====
fe_skills = [
    '前端开发：熟练 HTML5 / CSS3 / JavaScript，能独立开发单页应用，掌握响应式布局与现代CSS特性',
    'UI设计：熟悉 iOS 风格毛玻璃设计体系，注重交互细节与用户体验，能独立完成界面设计',
    'AI辅助开发：熟练运用 AI 工具进行代码生成、调试优化、需求拆解，开发效率提升显著',
    '后端基础：了解 Java 与 Spring Boot 基本原理，熟悉 RESTful API 设计，能对接后端接口',
    '工程化：熟练使用 Git 版本管理，有 GitHub 项目托管与 Pages 部署经验，具备完整上线能力',
    '存储技术：了解 localStorage、File System Access API、IndexedDB 等浏览器本地存储方案'
]

fe_project_desc = '独立开发的单文件HTML个人工作台应用，整合9大功能模块，纯前端实现无需服务器。全程使用AI辅助开发，从需求到上线效率大幅提升。已部署GitHub Pages，手机电脑均可访问。'

fe_works = [
    '前端架构：设计左侧导航+右侧内容区的单页应用架构，9个模块无刷新切换，状态管理清晰',
    'UI设计：采用iOS风格毛玻璃设计体系，半透明磨砂质感，响应式适配桌面与移动端',
    '功能开发：独立完成今日计划、自媒体追踪、项目管理、健身记录、饮食营养计算、攒钱目标等全部业务模块',
    'AI辅助：全程运用AI工具辅助开发，包括代码生成、Bug调试、样式优化、需求拆解',
    '数据持久化：实现本地数据自动保存与备份恢复，兼容File System Access API与localStorage双方案',
    '工程化：Git版本管理12次迭代提交，部署GitHub Pages生成在线链接'
]

fe_evals = [
    '前端开发能力扎实，能独立完成从界面设计到代码实现的全流程，注重用户体验与视觉细节',
    '熟练运用AI工具提升开发效率，善于借助AI解决技术问题，快速学习新技术',
    '有完整项目从0到1的独立开发经验，熟悉软件开发全流程',
    '学习能力强，做事踏实细心，对前端技术有热情，喜欢动手做实用的产品'
]

create_resume(
    r'C:\Users\a1773\Desktop\kk的APP\简历-全栈偏前端方向.docx',
    '求职意向：全栈开发工程师（偏前端） | 2026年6月毕业',
    fe_skills,
    fe_project_desc,
    fe_works,
    fe_evals
)
