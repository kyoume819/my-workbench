from docx import Document
from docx.shared import Pt, RGBColor, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn

doc = Document()

# 设置默认字体
doc.styles['Normal'].font.name = '微软雅黑'
doc.styles['Normal']._element.rPr.rFonts.set(qn('w:eastAsia'), '微软雅黑')
doc.styles['Normal'].font.size = Pt(11)

# ===== 标题 =====
title = doc.add_paragraph()
title.alignment = WD_ALIGN_PARAGRAPH.CENTER
run = title.add_run('个人简历')
run.font.size = Pt(22)
run.font.bold = True
run.font.color.rgb = RGBColor(0x1A, 0x1A, 0x2E)

# 基本信息
info = doc.add_paragraph()
info.alignment = WD_ALIGN_PARAGRAPH.CENTER
info_run = info.add_run('求职意向：项目经理助理 | 2026年6月毕业 | 河北科技学院 · 软件工程')
info_run.font.size = Pt(11)
info_run.font.color.rgb = RGBColor(0x55, 0x55, 0x55)

doc.add_paragraph()

# ===== 教育背景 =====
h2 = doc.add_paragraph()
h2_run = h2.add_run('教育背景')
h2_run.font.size = Pt(14)
h2_run.font.bold = True
h2_run.font.color.rgb = RGBColor(0x00, 0x7A, 0xFF)

edu = doc.add_paragraph()
edu.add_run('河北科技学院  |  软件工程（本科）  |  2022.09 - 2026.06').font.size = Pt(11)

doc.add_paragraph()

# ===== 专业技能 =====
h2 = doc.add_paragraph()
h2_run = h2.add_run('专业技能')
h2_run.font.size = Pt(14)
h2_run.font.bold = True
h2_run.font.color.rgb = RGBColor(0x00, 0x7A, 0xFF)

skills = [
    '项目管理：熟悉软件开发全流程，能独立编写PRD、开发计划、进度追踪文档',
    '任务管理：熟练使用任务管理工具，能拆解任务、跟进进度、协调资源',
    '文档能力：能独立撰写需求文档、项目总结、会议纪要等各类项目文档',
    '沟通协调：具备良好的沟通能力和团队协作意识，能对接开发、设计、测试等不同角色',
    '工具：熟练使用 Office 办公软件，了解 Git 版本管理，有项目部署上线经验',
    '技术背景：了解前端开发基本流程，能与开发团队有效沟通，理解技术术语和需求'
]
for s in skills:
    p = doc.add_paragraph(style='List Bullet')
    p.add_run(s).font.size = Pt(11)

doc.add_paragraph()

# ===== 项目经历（重点） =====
h2 = doc.add_paragraph()
h2_run = h2.add_run('项目经历')
h2_run.font.size = Pt(14)
h2_run.font.bold = True
h2_run.font.color.rgb = RGBColor(0x00, 0x7A, 0xFF)

# 项目标题
proj_title = doc.add_paragraph()
proj_title.add_run('个人工作台 — 个人工作生活管理工具').font.bold = True
proj_title.runs[0].font.size = Pt(12)

# 项目链接
link_p = doc.add_paragraph()
link_run = link_p.add_run('在线体验：https://kyoume819.github.io/my-workbench/')
link_run.font.size = Pt(10)
link_run.font.color.rgb = RGBColor(0x00, 0x66, 0xCC)

link_p2 = doc.add_paragraph()
link_run2 = link_p2.add_run('代码仓库：https://github.com/kyoume819/my-workbench')
link_run2.font.size = Pt(10)
link_run2.font.color.rgb = RGBColor(0x00, 0x66, 0xCC)

# 项目描述
desc = doc.add_paragraph()
desc.add_run('项目描述：').font.bold = True
desc.add_run('作为项目负责人，独立完成从需求分析、PRD撰写、开发计划制定到最终上线的全流程管理。项目包含9个功能模块，覆盖工作、自媒体、健身、饮食、财务等多个场景，已部署上线并持续迭代。')

# 项目管理工作
work = doc.add_paragraph()
work.add_run('项目管理工作：').font.bold = True

works = [
    '需求管理：独立撰写产品需求文档（PRD），明确产品目标、功能边界、验收标准，完成需求拆解与优先级排序',
    '进度管理：制定分阶段开发计划，将项目拆分为12个阶段任务，跟踪每个阶段完成情况，确保按计划交付',
    '文档管理：编写并维护PRD文档、开发计划、设计原型说明等各类项目文档，版本清晰、可追溯',
    '迭代管理：经历4轮需求迭代优化，每轮收集反馈、调整优先级、跟进落地，持续改进产品体验',
    '上线部署：负责项目版本管理与部署上线，使用Git进行版本控制，通过GitHub Pages发布，手机电脑均可访问',
    '质量把控：对照验收标准逐条验证功能，发现问题及时修复，确保交付质量'
]
for w in works:
    p = doc.add_paragraph(style='List Bullet')
    p.add_run(w).font.size = Pt(11)

doc.add_paragraph()

# ===== 自我评价 =====
h2 = doc.add_paragraph()
h2_run = h2.add_run('自我评价')
h2_run.font.size = Pt(14)
h2_run.font.bold = True
h2_run.font.color.rgb = RGBColor(0x00, 0x7A, 0xFF)

evals = [
    '有完整项目从0到1的管理经验，熟悉软件开发全流程，能快速理解需求并推进落地',
    '做事细心有条理，擅长文档撰写和进度跟踪，能把复杂任务拆解成可执行的小步骤',
    '沟通能力好，有技术背景，能和开发、设计、测试等不同角色有效对接',
    '执行力强，有责任心，善于协调资源解决问题，推动项目按计划完成'
]
for e in evals:
    p = doc.add_paragraph(style='List Bullet')
    p.add_run(e).font.size = Pt(11)

# 保存
doc.save(r'C:\Users\a1773\Desktop\kk的APP\简历-项目经理助理方向.docx')
print('简历生成完成')
