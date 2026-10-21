from docx import Document
from docx.shared import Pt, RGBColor, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn

doc = Document()

# 设置默认字体
doc.styles['Normal'].font.name = '微软雅黑'
doc.styles['Normal']._element.rPr.rFonts.set(qn('w:eastAsia'), '微软雅黑')
doc.styles['Normal'].font.size = Pt(11)

# ===== 标题 =====
title = doc.add_paragraph()
title.alignment = WD_ALIGN_PARAGRAPH.CENTER
run = title.add_run('个人简历')
run.font.size = Pt(22)
run.font.bold = True
run.font.color.rgb = RGBColor(0x1A, 0x1A, 0x2E)

# 基本信息
info = doc.add_paragraph()
info.alignment = WD_ALIGN_PARAGRAPH.CENTER
info_run = info.add_run('求职意向：全栈开发工程师 | 2026年6月毕业 | 河北科技学院 · 软件工程')
info_run.font.size = Pt(11)
info_run.font.color.rgb = RGBColor(0x55, 0x55, 0x55)

doc.add_paragraph()

# ===== 教育背景 =====
h2 = doc.add_paragraph()
h2_run = h2.add_run('教育背景')
h2_run.font.size = Pt(14)
h2_run.font.bold = True
h2_run.font.color.rgb = RGBColor(0x00, 0x7A, 0xFF)

edu = doc.add_paragraph()
edu.add_run('河北科技学院  |  软件工程（本科）  |  2022.09 - 2026.06').font.size = Pt(11)

doc.add_paragraph()

# ===== 专业技能 =====
h2 = doc.add_paragraph()
h2_run = h2.add_run('专业技能')
h2_run.font.size = Pt(14)
h2_run.font.bold = True
h2_run.font.color.rgb = RGBColor(0x00, 0x7A, 0xFF)

skills = [
    '前端：熟悉 HTML5 / CSS3 / JavaScript，能独立开发单页应用，掌握响应式布局与现代UI设计',
    '后端：了解 Java 基础，熟悉Spring Boot基本概念，了解RESTful API设计',
    '工具：熟练使用 Git 版本管理，有GitHub项目托管与Pages部署经验',
    '存储：了解 localStorage、File System Access API、IndexedDB 等浏览器存储技术',
    '设计：熟悉 iOS 风格毛玻璃UI设计，注重用户体验与界面细节',
    '其他：独立完成从需求分析、UI设计到代码实现的全流程开发能力'
]
for s in skills:
    p = doc.add_paragraph(style='List Bullet')
    p.add_run(s).font.size = Pt(11)

doc.add_paragraph()

# ===== 项目经历（重点） =====
h2 = doc.add_paragraph()
h2_run = h2.add_run('项目经历')
h2_run.font.size = Pt(14)
h2_run.font.bold = True
h2_run.font.color.rgb = RGBColor(0x00, 0x7A, 0xFF)

# 项目标题
proj_title = doc.add_paragraph()
proj_title.add_run('个人工作台 — 本地工作生活管理工具').font.bold = True
proj_title.runs[0].font.size = Pt(12)

# 项目链接
link_p = doc.add_paragraph()
link_run = link_p.add_run('在线体验：https://kyoume819.github.io/my-workbench/')
link_run.font.size = Pt(10)
link_run.font.color.rgb = RGBColor(0x00, 0x66, 0xCC)

link_p2 = doc.add_paragraph()
link_run2 = link_p2.add_run('代码仓库：https://github.com/kyoume819/my-workbench')
link_run2.font.size = Pt(10)
link_run2.font.color.rgb = RGBColor(0x00, 0x66, 0xCC)

# 项目描述
desc = doc.add_paragraph()
desc.add_run('项目描述：').font.bold = True
desc.add_run('独立设计开发的单文件HTML个人工作台应用，整合工作、自媒体、健身、饮食、攒钱等9大功能模块，纯前端实现，无需服务器和数据库，打开浏览器即可使用。已部署到GitHub Pages，支持手机和电脑访问。')

# 技术栈
tech = doc.add_paragraph()
tech.add_run('技术栈：').font.bold = True
tech.add_run('HTML5 / CSS3 / 原生JavaScript / File System Access API / localStorage / Git / GitHub Pages')

# 主要工作
work = doc.add_paragraph()
work.add_run('主要工作：').font.bold = True

works = [
    '架构设计：设计左侧导航+右侧内容区的单页应用架构，实现9个模块无刷新切换，路由状态管理清晰',
    '数据持久化：实现本地数据自动保存与备份恢复功能，使用File System Access API和localStorage双方案兼容，刷新关闭不丢失',
    '功能开发：独立完成今日计划、自媒体追踪、开发项目管理、咨询客户跟进、健身记录、饮食营养计算、攒钱目标等全部业务模块',
    'UI设计：采用iOS风格毛玻璃UI设计，半透明磨砂质感，响应式适配桌面和移动端，注重交互细节',
    '工程化：使用Git进行版本管理，12次迭代提交，部署到GitHub Pages生成在线访问链接，可直接在浏览器中体验',
    '全流程：从需求分析、PRD文档、开发计划、UI设计到代码实现全程独立完成，具备完整项目开发能力'
]
for w in works:
    p = doc.add_paragraph(style='List Bullet')
    p.add_run(w).font.size = Pt(11)

doc.add_paragraph()

# ===== 自我评价 =====
h2 = doc.add_paragraph()
h2_run = h2.add_run('自我评价')
h2_run.font.size = Pt(14)
h2_run.font.bold = True
h2_run.font.color.rgb = RGBColor(0x00, 0x7A, 0xFF)

evals = [
    '有独立完成完整项目的能力，从0到1独立开发并部署上线了个人工作台应用',
    '对全栈开发有热情，既做前端界面设计，也理解后端数据存储与API设计思路',
    '学习能力强，能快速上手新技术，喜欢动手做实用的工具解决实际问题',
    '做事踏实细心，注重用户体验和代码质量，项目功能完整、细节到位'
]
for e in evals:
    p = doc.add_paragraph(style='List Bullet')
    p.add_run(e).font.size = Pt(11)

# 保存
doc.save(r'C:\Users\a1773\Desktop\kk的APP\简历.docx')
print('简历Word生成完成')
